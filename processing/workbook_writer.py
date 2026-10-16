"""Replay the retained Excel builder's operations and cache verified metrics."""
import io,copy,json,zipfile,xml.etree.ElementTree as ET

def save_plan(plan,destination):
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule,FormulaRule
    from openpyxl.utils import get_column_letter as letter
    from openpyxl.worksheet.formula import ArrayFormula
    from openpyxl.workbook.properties import CalcProperties
    wb=openpyxl.Workbook();wb.remove(wb.active)
    sheets={name:wb.create_sheet(name) for name in plan['names']}
    freezes={name:[0,0] for name in sheets}
    hexcolor=lambda value:value.lstrip('#') if isinstance(value,str) else value
    def font_args(values):
        return {({'size':'sz','bold':'b','italic':'i'}.get(k,k)):hexcolor(v) if k=='color' else v for k,v in values.items() if k in ['name','size','bold','italic','color']}
    def side(values):return Side(style=values.get('style','thin'),color=hexcolor(values.get('color','#DDE3E8')))
    for operation in plan['operations']:
        ws=sheets[operation['sheet']];path=operation['path'];value=operation['value'];area=operation.get('range')
        if path==['tabColor']:ws.sheet_properties.tabColor=hexcolor(value);continue
        if path==['showGridLines']:ws.sheet_view.showGridLines=value;continue
        if path==['freezeRows']:freezes[ws.title][0]=value;continue
        if path==['freezeColumns']:freezes[ws.title][1]=value;continue
        if not area:continue
        row,col,height,width=area
        ref=f'{letter(col)}{row}:{letter(col+width-1)}{row+height-1}'
        if path[0] in ['values','formulas']:
            for i,values in enumerate(value):
                for j,item in enumerate(values):ws.cell(row+i,col+j,item)
            continue
        if path==['merge']:ws.merge_cells(ref);continue
        if path==['dataValidation']:
            rule=value['rule'];formula1=rule.get('formula1');formula2=rule.get('formula2')
            operator={'equal':'equal','greaterThanOrEqual':'greaterThanOrEqual'}.get(rule.get('operator'),rule.get('operator'))
            if rule['type']=='list':formula1='"'+','.join(str(x) for x in rule.get('values',[]))+'"'
            dv=DataValidation(type=rule['type'],operator=operator,formula1=formula1,formula2=formula2,allow_blank=True)
            ws.add_data_validation(dv);dv.add(ref);continue
        if path[0]=='conditional':
            fmt=value.get('format',{});fill=PatternFill('solid',fgColor=hexcolor(fmt['fill'])) if fmt.get('fill') else None
            font=Font(**font_args(fmt.get('font',{}))) if fmt.get('font') else None
            if path[1]=='cellIs':rule=CellIsRule(operator=value['operator'],formula=[str(value.get('formula',0))],fill=fill,font=font)
            elif path[1]=='custom':rule=FormulaRule(formula=[value['formula']],fill=fill,font=font)
            elif path[1]=='containsText':rule=FormulaRule(formula=[f'ISNUMBER(SEARCH("{value["text"]}",{letter(col)}{row}))'],fill=fill,font=font)
            else:continue
            ws.conditional_formatting.add(ref,rule);continue
        if path==['format','columnWidth']:
            for c in range(col,col+width):ws.column_dimensions[letter(c)].width=value
            continue
        if path==['format','rowHeight']:
            for r in range(row,row+height):ws.row_dimensions[r].height=value
            continue
        for rows in ws.iter_rows(min_row=row,max_row=row+height-1,min_col=col,max_col=col+width-1):
            for cell in rows:
                if path==['numberFormat']:cell.number_format=value
                elif path==['format','fill']:cell.fill=PatternFill('solid',fgColor=hexcolor(value))
                elif path[:2]==['format','font']:
                    current=copy.copy(cell.font)
                    args=font_args(value if len(path)==2 else {path[2]:value})
                    for key,setting in args.items():setattr(current,key,setting)
                    cell.font=current
                elif path[:2]==['format','borders']:
                    current=copy.copy(cell.border)
                    if value.get('preset')=='all':
                        edge=side(value)
                        for key in ['left','right','top','bottom']:setattr(current,key,edge)
                    for key,settings in value.items():
                        if key in ['left','right','top','bottom']:setattr(current,key,side(settings))
                        if key=='insideHorizontal':
                            if cell.row>row:current.top=side(settings)
                            if cell.row<row+height-1:current.bottom=side(settings)
                        if key=='insideVertical':
                            if cell.column>col:current.left=side(settings)
                            if cell.column<col+width-1:current.right=side(settings)
                    cell.border=current
                elif path[:1]==['format'] and len(path)==2:
                    attr={'horizontalAlignment':'horizontal','verticalAlignment':'vertical','wrapText':'wrap_text'}.get(path[1])
                    if attr:
                        alignment=copy.copy(cell.alignment)
                        if attr=='vertical' and value=='center':value='center'
                        setattr(alignment,attr,value);cell.alignment=alignment
    for name,(row,col) in freezes.items():
        if row or col:sheets[name].freeze_panes=f'{letter(col+1)}{row+1}'
    # MINVERSE is the original full-BOM array, not a single scalar formula.
    array=plan['array'];ws=sheets[array['sheet']];r,c,h,w=[array[k] for k in ['row','col','height','width']]
    formula=ws.cell(r,c).value;ws.cell(r,c).value=ArrayFormula(ref=f'{letter(c)}{r}:{letter(c+w-1)}{r+h-1}',text=formula)
    for cell,value in plan['caches'][ws.title].items():
        target=ws[cell]
        if target.value is None:target.value=value
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    stream=io.BytesIO();wb.save(stream)
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    ET.register_namespace('',ns['m'])
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source,zipfile.ZipFile(destination,'w',zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            raw=source.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
                index=int(item.filename.rsplit('sheet',1)[1].split('.')[0])-1
                cache=plan['caches'].get(plan['names'][index],{})
                document=ET.fromstring(raw)
                for element in document.findall('.//m:sheetData/m:row/m:c',ns):
                    key=element.attrib.get('r')
                    if key not in cache or element.find('m:f',ns) is None:continue
                    value=cache[key];existing=element.find('m:v',ns)
                    if existing is None:existing=ET.SubElement(element,'{'+ns['m']+'}v')
                    if isinstance(value,bool):element.set('t','b');existing.text='1' if value else '0'
                    elif isinstance(value,(int,float)):element.attrib.pop('t',None);existing.text=str(value)
                    elif value is None:element.attrib.pop('t',None);existing.text=None
                    else:element.set('t','str');existing.text=str(value)
                raw=ET.tostring(document,encoding='utf-8',xml_declaration=True)
            target.writestr(item,raw)
