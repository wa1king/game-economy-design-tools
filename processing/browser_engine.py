import json,math,io,base64
from copy import deepcopy
import scheduler
from js import report_progress
_bootstrap=json.load(open('/home/pyodide/bootstrap.json'))
SEED=_bootstrap['design']
SCHEMA=_bootstrap['schema']
def validate(payload, allow_draft=False):
    source = payload.get('design',payload)
    if not isinstance(source,dict) or not isinstance(source.get('products'),list):
        raise ValueError('方案须包含商品列表。')
    buildings=source.get('buildings',SEED['buildings']); lands=source.get('lands_by_level',SEED['lands_by_level'])
    if not isinstance(buildings,list) or not isinstance(lands,dict):raise ValueError('建筑或土地配置格式不正确。')
    building_map={}; normal_buildings=[]
    for b in buildings:
        if not isinstance(b,dict):raise ValueError('建筑配置格式不正确。')
        bid=b.get('id'); name=b.get('name'); level=b.get('level'); color=b.get('color','#DDEBF7')
        if not isinstance(bid,int) or isinstance(bid,bool) or bid<=0 or bid in building_map:raise ValueError('建筑ID需为不重复的正整数。')
        if not isinstance(name,str) or not name.strip() or len(name)>50:raise ValueError('请填写建筑名称，最多50字。')
        if not isinstance(level,int) or isinstance(level,bool) or not 1<=level<=20:raise ValueError(f'{name}：建筑解锁等级需为1—20。')
        import re
        if not isinstance(color,str) or not re.fullmatch(r'#[0-9a-fA-F]{6}',color):raise ValueError(f'{name}：颜色格式不正确。')
        normalized={'id':bid,'name':name.strip(),'level':level,'color':color}; building_map[bid]=normalized;normal_buildings.append(normalized)
    normal_lands={}; previous=0
    for level in range(1,21):
        amount=lands.get(str(level),lands.get(level))
        if isinstance(amount,bool) or not isinstance(amount,int) or amount<1 or amount<previous:raise ValueError('各级累计土地需为正整数，且不能减少。')
        if level==1 and amount!=6:raise ValueError('初始土地为6片。')
        if level>1 and amount-previous not in (0,6):raise ValueError('每级新增土地只能为0或6片。')
        normal_lands[str(level)]=amount;previous=amount
    found, products = set(), []
    for item in source['products']:
        if not isinstance(item,dict): raise ValueError('商品格式不正确。')
        key = item.get('id')
        if isinstance(key,bool) or not isinstance(key,int) or key<1 or key in found:raise ValueError('商品ID须为不重复的正整数。')
        found.add(key);name=item.get('name');level=item.get('level');building=item.get('building')
        if not isinstance(name,str) or not name.strip() or len(name)>50:raise ValueError('请填写商品名称，最多50字。')
        if isinstance(level,bool) or not isinstance(level,int) or not 1<=level<=20:raise ValueError(f'{name}：商品等级需为1—20。')
        if isinstance(building,bool) or not isinstance(building,int) or building!=0 and building not in building_map:raise ValueError(f'{name}：所属建筑不存在。')
        if not allow_draft and building and building_map[building]['level']>level:raise ValueError(f'{name}：所属建筑尚未解锁。')
        original={'id':key,'name':name.strip(),'level':level,'building':building,'buildingName':building_map[building]['name'] if building else '农田','type':'加工品' if building else '作物'}
        if item.get('yield',1) != 1: raise ValueError(f'{name}：单次产量固定为1。')
        duration=item.get('time')
        if isinstance(duration,bool) or not isinstance(duration,(int,float)) or not math.isfinite(duration) or duration<=0:
            raise ValueError(f'{name}：单步时间必须为正数。')
        inputs=item.get('inputs',[])
        if not isinstance(inputs,list): raise ValueError(f'{name}：配方格式不正确。')
        used=set(); recipe=[]
        for pair in inputs:
            if not isinstance(pair,list) or len(pair)!=2: raise ValueError(f'{name}：原料需包含ID和数量。')
            up,quantity=pair
            if isinstance(up,bool) or not isinstance(up,int) or up in used:
                raise ValueError(f'{name}：原料不存在或重复。')
            if up==key: raise ValueError(f'{name}：不能投入自身。')
            if isinstance(quantity,bool) or not isinstance(quantity,(int,float)) or not math.isfinite(quantity) or quantity<=0 or int(quantity)!=quantity:
                raise ValueError(f'{name}：投入数量必须为正整数。')
            used.add(up); recipe.append([up,int(quantity)])
        if not original['building'] and recipe: raise ValueError(f'{name}：作物不能填加工投入。')
        if not allow_draft and original['building'] and not recipe: raise ValueError(f'{name}：加工品至少需要一种原料。')
        products.append(dict(original,time=duration,inputs=sorted(recipe),**{'yield':1}))
    if len({b['name'] for b in normal_buildings})!=len(normal_buildings):raise ValueError('建筑名称不能重复。')
    if len({p['name'] for p in products})!=len(products):raise ValueError('商品名称不能重复。')
    normal_buildings.sort(key=lambda b:(b['level'],b['id']))
    order={0:0,**{b['id']:i+1 for i,b in enumerate(normal_buildings)}}
    products.sort(key=lambda p:(order[p['building']],p['level'],p['id'])); indexed={p['id']:p for p in products}; done={}; stack=[]
    for p in products:
        for up,q in p['inputs']:
            if up not in indexed:raise ValueError(f'{p["name"]}：原料不存在。')
            if not allow_draft and indexed[up]['level']>p['level']:raise ValueError(f'{p["name"]}：原料{indexed[up]["name"]}尚未解锁。')
    def count(key):
        if key in stack:
            cycle=stack[stack.index(key):]+[key]
            raise ValueError('配方循环：'+' → '.join(indexed[i]['name'] for i in cycle))
        if key not in done:
            stack.append(key); total=1+sum(q*count(i) for i,q in indexed[key]['inputs']); stack.pop()
            if total>500: raise ValueError(f'{indexed[key]["name"]}：全链超过500次生产，请缩小投入数量。')
            done[key]=total
        return done[key]
    for key in indexed: count(key)
    result={'products':products,'buildings':normal_buildings,'lands_by_level':normal_lands}
    if 'turnover' in source:
        result['turnover']=validate_turnover(source['turnover'],normal_buildings)
    if 'pricing' in source and source['pricing'] is None:raise ValueError('价格配置格式不正确。')
    result['pricing']=validate_pricing(source.get('pricing'),normal_buildings)
    return result

def validate_pricing(config,buildings):
    defaults=[(120,1),(300,2),(600,3),(1200,4),(1800,5),(3600,6),(7200,7),(10800,8),(14400,9),(18000,10)]
    if config is None:config={'version':1,'timeValues':[{'seconds':s,'value':v} for s,v in defaults]}
    if not isinstance(config,dict):raise ValueError('价格配置格式不正确。')
    if isinstance(config.get('version',1),bool) or config.get('version',1)!=1:raise ValueError('不支持此价格配置版本。')
    points=config.get('timeValues')
    coefficients=config.get('sourceCoefficients',{})
    if not isinstance(points,list) or len(points)<2 or not isinstance(coefficients,dict):raise ValueError('时间价值须配置至少2个时间点。')
    normal=[];seen=set()
    for point in points:
        if not isinstance(point,dict):raise ValueError('时间价值格式不正确。')
        seconds=point.get('seconds');value=point.get('value')
        if any(isinstance(n,bool) or not isinstance(n,(int,float)) or not math.isfinite(n) for n in (seconds,value)) or seconds<=0 or value<0:raise ValueError('时间须为正数，价值须为非负数。')
        if seconds in seen:raise ValueError('时间价值中存在重复时间。')
        seen.add(seconds);normal.append({'seconds':seconds,'value':value})
    normal.sort(key=lambda p:p['seconds'])
    for previous,point in zip(normal,normal[1:]):
        if point['seconds']-previous['seconds']<=1e-6:raise ValueError('时间点不能重复。')
        if point['value']<previous['value']:raise ValueError('时间越长，价值不能降低。')
    if any(isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or value<0 for value in coefficients.values()):raise ValueError('来源系数须为非负数。')
    factors={}
    for bid in [0]+[b['id'] for b in buildings]:
        factor=coefficients.get(str(bid),coefficients.get(bid,1 if bid==0 else 2))
        if isinstance(factor,bool) or not isinstance(factor,(int,float)) or not math.isfinite(factor) or factor<0:raise ValueError('来源价值系数须为非负数。')
        factors[str(bid)]=factor
    return {'version':1,'timeValues':normal,'sourceCoefficients':factors}

def validate_turnover(config,buildings):
    if not isinstance(config,dict):raise ValueError('队列设置格式不正确。')
    active=config.get('includeActive',True);queues=config.get('queues',{})
    if not isinstance(active,bool) or not isinstance(queues,dict):raise ValueError('队列设置格式不正确。')
    normal={}
    for b in buildings:
        values=queues.get(str(b['id']),[2]*20)
        if not isinstance(values,list) or len(values)!=20 or any(isinstance(n,bool) or not isinstance(n,int) or not 0<=n<=9007199254740991 for n in values):raise ValueError(b['name']+'：队列等待长度须为20级的非负整数。')
        normal[str(b['id'])]=values[:]
    return {'version':1,'includeActive':active,'queues':normal}

def read_excel(data):
    import openpyxl
    wb=openpyxl.load_workbook(io.BytesIO(data),data_only=False,read_only=True)
    try:
        if SCHEMA['sheet'] not in wb.sheetnames: raise ValueError('未找到01_数值调参，请导入本工具导出的Excel。')
        ws=wb[SCHEMA['sheet']]; products=[]; buildings=deepcopy(SEED['buildings']);lands=deepcopy(SEED['lands_by_level'])
        if '00_等级规划' in wb.sheetnames:
            lv=wb['00_等级规划']; cumulative=0;lands={}
            for level in range(1,21):
                addition=lv.cell(level+4,2).value
                if not isinstance(addition,(int,float)) or isinstance(addition,bool) or int(addition)!=addition:raise ValueError('新增土地需为整数。')
                cumulative+=int(addition);lands[str(level)]=cumulative
            buildings=[];row=29
            while lv.cell(row,1).value is not None:
                buildings.append({'id':lv.cell(row,1).value,'name':lv.cell(row,2).value,'level':lv.cell(row,3).value,'color':lv.cell(row,4).value});row+=1
        by_name={b['name']:b for b in buildings}; fixed=[];c=5
        while ws.cell(5,c).value is not None:
            bid=0 if ws.cell(8,c).value=='农田' else by_name.get(ws.cell(8,c).value,{}).get('id')
            if bid is None:raise ValueError('商品所属建筑不在建筑配置中。')
            fixed.append({'id':ws.cell(5,c).value,'name':ws.cell(6,c).value,'level':ws.cell(7,c).value,'building':bid});c+=1
        if not fixed:raise ValueError('01表未配置商品。')
        for j,p in enumerate(fixed):
            c=5+j
            inputs=[]
            for i,up in enumerate(fixed):
                q=ws.cell(18+i,c).value
                if q is None:q=0
                if not isinstance(q,(int,float)) or isinstance(q,bool) or not math.isfinite(q) or int(q)!=q:
                    raise ValueError(f'{p["name"]}：配方必须为整数。')
                if i==j:
                    if q!=1:raise ValueError(f'{p["name"]}：单次产量固定为1。')
                elif q>0:raise ValueError(f'{p["name"]}：Excel配方投入须填负数。')
                elif q:inputs.append([up['id'],-q])
            products.append(dict(p,time=ws.cell(9,c).value,inputs=inputs))
        source={'products':products,'buildings':buildings,'lands_by_level':lands}
        if '05_周转容量' in wb.sheetnames:
            queue=wb['05_周转容量'];active=queue.cell(3,3).value
            if active not in (0,1) or isinstance(active,bool):raise ValueError('05表的在产位开关须为0或1。')
            settings={'includeActive':bool(active),'queues':{}}
            for col in range(3,3+len(buildings)):
                key=queue.cell(6,col).value
                if not isinstance(key,str) or not key.startswith('建筑:'):raise ValueError('05表建筑ID格式不正确。')
                values=[]
                for level in range(1,21):
                    value=queue.cell(level+6,col).value
                    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or value<0 or int(value)!=value:raise ValueError('05表队列等待长度须为非负整数。')
                    values.append(int(value))
                bid=key.split(':',1)[1]
                if bid in settings['queues']:raise ValueError('05表建筑ID重复。')
                settings['queues'][bid]=values
            if set(settings['queues'])!={str(b['id']) for b in buildings}:raise ValueError('05表建筑ID与等级规划不一致。')
            source['turnover']=settings
        if '06_价格配置' in wb.sheetnames:
            price=wb['06_价格配置'];points=[]
            if price.max_row is None:price.calculate_dimension(force=True)
            header=next((row for row in range(5,price.max_row+1) if price.cell(row,1).value=='商品ID'),None)
            if header is None:raise ValueError('价格配置中未找到商品计算表头。')
            for row in range(5,header):
                seconds=price.cell(row,1).value;value=price.cell(row,2).value
                if seconds is None and value is None:continue
                if not isinstance(seconds,(int,float)) or isinstance(seconds,bool):raise ValueError('价格配置中的时间须为数字。')
                points.append({'seconds':seconds,'value':value})
            factors={}
            for row in range(5,header):
                key=price.cell(row,5).value
                if key is None and price.cell(row,7).value is None:continue
                if not isinstance(key,str) or not key.startswith('来源:'):raise ValueError('价格配置中的来源ID格式不正确。')
                bid=key.split(':',1)[1]
                if bid in factors:raise ValueError('价格配置中来源ID重复。')
                factors[bid]=price.cell(row,7).value
            expected={'0'}|{str(b['id']) for b in buildings}
            if set(factors)!=expected:raise ValueError('价格配置中的来源ID与当前建筑不一致。')
            source['pricing']={'version':1,'timeValues':points,'sourceCoefficients':factors}
        return validate(source)
    finally:wb.close()

def execute(kind,payload):
    if kind=='import':return {'design':read_excel(base64.b64decode(payload['base64']))}
    design=validate(payload,allow_draft=kind=='validate')
    if kind=='validate':return {'design':design}
    result=scheduler.analyze(design['products'],design['lands_by_level'],time_limit_seconds=2,node_limit=20000,progress_callback=lambda i,n,name:report_progress(i,n,name))
    return {'design':design,'analysis':result}
